"""Tests for tamarket.mhclg on synthetic fixtures (see conftest.py)."""

from __future__ import annotations

import warnings

import numpy as np
import openpyxl
import pandas as pd
import pytest

from conftest import write_book
from tamarket import mhclg
from tamarket.mhclg import (
    ACCOMMODATION_TYPES,
    TA_HOUSEHOLDS_COLUMNS,
    TA_SPEND_COLUMNS,
    TaMarketWarning,
    discover_raw_files,
    financial_year_of_quarter,
    load_ta_households,
    load_ta_spend,
    quarter_from_text,
    read_ro_metadata,
    read_ro_timeseries,
    read_ta_households,
    read_ta_spend,
    total_ta_spend,
    validate_ta_households,
    year_from_text,
)


def value(df, code, key, column="accommodation_type", field="households"):
    hit = df[(df["la_code"] == code) & (df[column] == key)][field]
    assert len(hit) == 1, (code, key, len(hit))
    return hit.iloc[0]


# ---------------------------------------------------------------- periods

@pytest.mark.parametrize("name, quarter", [
    ("Statutory_Homelessness_Detailed_Local_Authority_Data_202603.ods", "2026-Q1"),
    ("Statutory_Homelessness_Detailed_Local_Authority_Data_202512_revised.ods", "2025-Q4"),
    ("Detailed_LA_202506_revised_corrected.xlsx", "2025-Q2"),
    ("V2_Detailed_LA_202209_revised.ods", "2022-Q3"),
    ("DetailedLA_202109_revised.ods", "2021-Q3"),
    ("Statutory_Homelessness_Detailed_Local_Authority_Data_2024-2025_corrected.ods", "2025-Q1"),
    ("Detailed_LA_20222023_revised_corrected.ods", "2023-Q1"),
    ("DetailedLA_2019-20__revised__Revised_dropdowns_fixed.ods", "2020-Q1"),
    ("Detailed_LA_2021-22__Revised_Nov_2023_.ods", "2022-Q1"),
    ("Households in temporary accommodation at 30 September 2025", "2025-Q3"),
    ("Detailed local authority level tables: October to December 2025", "2025-Q4"),
    ("no date here", None),
])
def test_quarter_from_text(name, quarter):
    assert quarter_from_text(name) == quarter


@pytest.mark.parametrize("name, year", [
    ("RO4_LA_Data_2024-25_data_by_LA.ods", "2024-25"),
    ("RO4_2023-24_data_by_LA.ods", "2023-24"),
    ("RO4_2020-21_data_by_LA_March_2025.ods", "2020-21"),
    ("RO4_2021-22_data_by_LA_Live.ods", "2021-22"),
    ("Revenue outturn 2019 to 2020", None),
    ("2024-2025", "2024-25"),
])
def test_year_from_text(name, year):
    assert year_from_text(name) == year


def test_cell_conversion():
    assert mhclg._to_number("\u00a31,000") == 1000
    assert mhclg._to_number("(5)") == -5
    assert mhclg._to_number(" 1,234.5 ") == 1234.5
    for marker in ["..", "[x]", "[c]", "-", "\u2013", "Not set", ""]:
        assert np.isnan(mhclg._to_number(marker)), marker
    assert mhclg._text("Bed\u00a0and  breakfast ") == "Bed and breakfast"
    assert mhclg._text(80.0) == "80"
    assert mhclg._norm("Barking & Dagenham [note 3]") == "barking and dagenham"


def test_financial_year_of_quarter():
    assert financial_year_of_quarter("2025-Q2") == "2025-26"
    assert financial_year_of_quarter("2025-Q4") == "2025-26"
    assert financial_year_of_quarter("2026-Q1") == "2025-26"
    assert financial_year_of_quarter("2025-Q1") == "2024-25"
    with pytest.raises(ValueError):
        financial_year_of_quarter("2025 Q1")


# ---------------------------------------------------------------- households (TA1)

def test_legacy_ta1_layout(legacy_ta1):
    with warnings.catch_warnings():
        warnings.simplefilter("error", TaMarketWarning)
        df = read_ta_households(legacy_ta1)
    assert list(df.columns) == TA_HOUSEHOLDS_COLUMNS
    assert set(df["la_code"]) == {"E09000001", "E09000002", "E07000223", "E06000001"}
    assert set(df["quarter"]) == {"2026-Q1"}
    assert set(df["accommodation_type"]) == set(ACCOMMODATION_TYPES) | {"total"}
    assert value(df, "E09000002", "nightly_paid_self_contained") == 900
    assert value(df, "E09000002", "bed_and_breakfast") == 150  # not the 'of which' columns
    assert value(df, "E09000002", "total") == 2000
    assert value(df, "E06000001", "total") == 1050  # '1,050' read as a number
    assert np.isnan(value(df, "E07000223", "hostels"))  # '..' is missing, not zero
    assert value(df, "E09000002", "hostels") == 50
    names = dict(zip(df["la_code"], df["la_name"]))
    assert names["E09000002"] == "Borough B"  # footnote marker removed
    assert df.attrs["column_map"]["total"].startswith("Total number of households in TA")
    assert df.attrs["non_numeric_values"][".."] == 7  # seven mapped columns for District C


def test_legacy_ta1_aggregates_optional(legacy_ta1):
    df = read_ta_households(legacy_ta1, include_aggregates=True)
    assert "E12000007" in set(df["la_code"])  # coded region kept
    assert "ENGLAND" not in set(df["la_name"])  # uncoded England row never becomes data


def test_validate_ta_households(legacy_ta1):
    check = validate_ta_households(read_ta_households(legacy_ta1)).set_index("la_code")
    assert check.loc[["E09000001", "E09000002", "E06000001"], "ok"].all()
    assert not check.loc["E07000223", "ok"]  # all missing


def test_accessible_single_header_layout(tmp_path):
    rows = [
        ["TA1: Households in temporary accommodation at 30 June 2025, by type of accommodation"],
        ["This worksheet contains one table. Some cells refer to notes in the notes worksheet."],
        ["Area code", "Area name", "Region name", "Total households in TA [note 1]",
         "Bed and breakfast hotels (including shared annexes)",
         "Bed and breakfast: households with children",
         "Nightly paid, privately managed accommodation, self-contained",
         "Hostels (including reception centres, emergency units and refuges)",
         "Private sector accommodation leased by the authority or by a registered provider",
         "Local authority or housing association (LA/HA) stock",
         "Any other type of accommodation",
         "Households in TA in another local authority district"],
        ["E92000001", "England", "", 600, 60, 30, 300, 20, 100, 70, 50, 90],
        ["E09000002", "Borough B", "Region L", 300, 30, 15, 150, 10, 50, 35, 25, 80],
        ["E08000035", "Metropolitan F", "Region Y", 280, 28, 14, 140, 9, 48, 34, 21, 5],
        ["E07000223", "District C", "Region S", "[x]", "[x]", "[x]", "[x]", "[x]", "[x]", "[x]", "[x]", "[x]"],
    ]
    path = write_book(tmp_path / "Detailed_LA_202506_revised_corrected.xlsx",
                      {"Contents": [["Contents"]], "TA1": rows})
    df = read_ta_households(path)
    assert set(df["quarter"]) == {"2025-Q2"}
    assert set(df["la_code"]) == {"E09000002", "E08000035", "E07000223"}
    assert value(df, "E08000035", "bed_and_breakfast") == 28
    assert value(df, "E08000035", "la_ha_stock") == 34
    assert value(df, "E09000002", "other") == 25
    assert set(df["la_name"]) == {"Borough B", "Metropolitan F", "District C"}  # not the region column
    assert validate_ta_households(df).set_index("la_code").loc[["E09000002", "E08000035"], "ok"].all()


def test_real_merged_cells_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TA1"
    ws.append(["Households in temporary accommodation, 31 March 2024"])
    ws.append([])
    ws.append([None, None, "Total number of households in TA", "Bed and breakfast hotels (including shared annexes)",
               None, "Nightly paid, privately managed accommodation, self-contained", None, "Hostels", None,
               "Private sector accommodation leased by your authority", None, "LA/HA stock", None,
               "Any other type of accommodation", None])
    ws.append(["Current ONS code", "Local authority name", None] + ["Total", "of which: households with children"] * 6)
    ws.append(["E06000001", "Unitary D", 100, 10, 5, 50, 20, 5, 0, 20, 10, 10, 5, 5, 0])
    ws.append(["E06000002", "Unitary G", 60, 0, 0, 30, 10, 0, 0, 20, 5, 5, 0, 5, 1])
    for start in range(4, 16, 2):  # D3:E3, F3:G3, ... as merged ranges
        ws.merge_cells(start_row=3, start_column=start, end_row=3, end_column=start + 1)
    ws.merge_cells("C3:C4")
    path = tmp_path / "Detailed_LA_202403.xlsx"
    wb.save(path)
    df = read_ta_households(path)
    assert set(df["quarter"]) == {"2024-Q1"}
    assert value(df, "E06000001", "nightly_paid_self_contained") == 50
    assert value(df, "E06000002", "private_sector_leased") == 20
    assert value(df, "E06000001", "total") == 100
    assert validate_ta_households(df)["ok"].all()


def test_of_which_super_header_and_title_quarter(tmp_path):
    rows = [
        ["Households in TA at 31 December 2025"],
        [],
        [None, None, "Total in TA", "of which:"],
        ["Code", "Local authority", None, "Bed and breakfast", "Nightly paid self-contained", "Hostels",
         "Private sector leased", "LA/HA stock", "Other"],
        ["E09000001", "Borough A", 100, 10, 40, 5, 20, 15, 10],
    ]
    df = read_ta_households(write_book(tmp_path / "ta_snapshot.ods", {"TA1": rows}))
    assert set(df["quarter"]) == {"2025-Q4"}
    assert value(df, "E09000001", "bed_and_breakfast") == 10
    assert value(df, "E09000001", "other") == 10
    assert validate_ta_households(df)["ok"].all()


def test_children_block_is_not_counted_twice(tmp_path):
    rows = [
        [None, None, "All households", None, None, "Households with children", None, None],
        ["Code", "Local authority", "Total", "Bed and breakfast", "Nightly paid", "Total", "Bed and breakfast",
         "Nightly paid"],
        ["E09000001", "Borough A", 50, 10, 40, 30, 2, 28],
    ]
    path = write_book(tmp_path / "Detailed_LA_202312_Revised_No_Dropdowns.ods", {"TA1": rows})
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        df = read_ta_households(path)
    assert not [w for w in caught if "columns look like" in str(w.message)]
    assert value(df, "E09000001", "bed_and_breakfast") == 10
    assert value(df, "E09000001", "total") == 50


def test_missing_sheet_and_missing_codes(tmp_path):
    path = write_book(tmp_path / "Detailed_LA_202303_Revised.ods", {"TA2": [["x"]]})
    with pytest.raises(ValueError, match="no sheet called 'TA1'"):
        read_ta_households(path)
    path = write_book(tmp_path / "Detailed_LA_202212_revised.ods", {"TA1": [["Total"], [1]]})
    with pytest.raises(ValueError, match="no ONS local authority codes"):
        read_ta_households(path)


def test_load_ta_households_prefers_quarterly_file(tmp_path, legacy_ta1):
    raw = legacy_ta1.parent
    annual = raw / "Statutory_Homelessness_Detailed_Local_Authority_Data_2025-2026.ods"
    annual.write_bytes(legacy_ta1.read_bytes())  # same 31 March 2026 snapshot
    with pytest.warns(TaMarketWarning, match="skipping"):
        df = load_ta_households(discover_raw_files(raw)["households"])
    assert set(df["source_file"]) == {legacy_ta1.name}


# ---------------------------------------------------------------- spend (RO4)

def test_ro4_wide_layout(ro4_wide):
    df = read_ta_spend(ro4_wide)
    assert list(df.columns) == TA_SPEND_COLUMNS
    assert set(df["year"]) == {"2024-25"}
    assert set(df["la_code"]) == {"E09000002", "E06000001", "E07000223"}  # no England, no county
    assert set(df["line"]) == set(mhclg.SPEND_LINES)
    assert value(df, "E09000002", "nightly_paid_self_contained", "line", "gbp") == 1_000_000
    assert value(df, "E06000001", "ta_admin", "line", "gbp") == 150_000
    assert np.isnan(value(df, "E07000223", "hostels", "line", "gbp"))  # 'Not set'
    totals = total_ta_spend(df).set_index("la_code")["gbp"]
    assert totals["E09000002"] == 1_950_000  # lines 80 to 86, not 87 or other council property
    assert np.isnan(totals["E07000223"])
    no_admin = total_ta_spend(df, include_admin=False).set_index("la_code")["gbp"]
    assert no_admin["E09000002"] == 1_875_000
    assert df.attrs["column_map"]["hostels"][0].endswith("Total expenditure")


def test_ro4_net_measure_and_aggregates(ro4_wide):
    df = read_ta_spend(ro4_wide, measure="net_current_expenditure", include_aggregates=True)
    assert {"E92000001", "E10000099"} <= set(df["la_code"])
    assert set(df["measure"]) == {"net_current_expenditure"}
    net = total_ta_spend(df).set_index("la_code")["gbp"]
    assert net["E09000002"] == 1_155_000
    assert net["E92000001"] == 115_500_000
    with pytest.raises(ValueError, match="measure must be one of"):
        read_ta_spend(ro4_wide, measure="gross")


def test_ro4_old_lines_are_combined(tmp_path):
    header = ["E-code", "ONS code", "Local authority",
              "Other nightly paid, privately managed accommodation - Total Expenditure",
              "Private managed accommodation leased by the authority - Total Expenditure",
              "Private managed accommodation leased by RSLs - Total Expenditure",
              "Bed and breakfast - Total Expenditure",
              "Homelessness: Administration - Total Expenditure"]
    rows = [["RO4 2019-20"], [], header, ["E0701", "E06000001", "Unitary D", 300, 70, 30, 90, 55]]
    path = write_book(tmp_path / "RO4_2019-20_data_by_LA.ods", {"RO4 LA Data": rows})
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        df = read_ta_spend(path)
    assert not [w for w in caught if "adding them" in str(w.message)]  # expected before 2020-21
    assert value(df, "E06000001", "private_sector_leased", "line", "gbp") == 100_000
    assert value(df, "E06000001", "nightly_paid_self_contained", "line", "gbp") == 300_000
    assert "ta_admin" not in set(df["line"])  # old homelessness administration is not TA admin


def test_load_ta_spend_two_years(tmp_path, ro4_wide):
    older = tmp_path / "RO4_2023-24_data_by_LA.ods"
    older.write_bytes(ro4_wide.read_bytes())
    df = load_ta_spend(discover_raw_files(tmp_path)["spend"])
    assert set(df["year"]) == {"2023-24", "2024-25"}
    assert len(df) == 2 * 3 * len(mhclg.SPEND_LINES)


# ---------------------------------------------------------------- multi-year data set

def test_ro_timeseries_with_metadata(tmp_path):
    dictionary = [
        ["Table 1: Data dictionary"],
        [],
        ["Data item name", "Data type", "Description", "Start date", "End date", "Form name", "Row number",
         "Column number", "Notes"],
        ["RS_hous_net_exp", "number", "Housing services (GFRA only): net current expenditure", "2017-18", "",
         "RS", 1, 1, ""],
        ["RO4_tanps_tot_exp", "number",
         "Nightly paid, privately managed accommodation, self-contained: Total expenditure",
         "2020-21", "", "RO4", 80, 3, ""],
        ["RO4_tanps_net_exp", "number",
         "Nightly paid, privately managed accommodation, self-contained: Net current expenditure",
         "2020-21", "", "RO4", 80, 7, ""],
        ["RO4_tabb_tot_exp", "number", "Bed and breakfast hotels (including shared annexes): Total expenditure",
         "2020-21", "", "RO4", 83, 3, ""],
        ["RO4_taadm_tot_exp", "number", "Temporary accommodation administration: Total expenditure",
         "2020-21", "", "RO4", 86, 3, ""],
        ["RO4_hra_tot_exp", "number",
         "Homeless Reduction Act: Administration, Prevention, Relief & Support: Total expenditure",
         "2020-21", "", "RO4", 87, 3, ""],
    ]
    meta = write_book(tmp_path / "Revenue_Outturn_metadata.ods",
                      {"Release log": [["Release log"], ["v4", "September 2026"]], "Data dictionary": dictionary})
    items = read_ro_metadata(meta)
    assert {"item", "form", "row", "column", "description"} <= set(items.columns)
    csv = tmp_path / "Revenue_Outturn_time_series_data_v4.csv"
    pd.DataFrame([
        ["202503", "E06", "UA", "Unitary Authorities total", "Not set", "Unitary Authority", "Unitary Authority",
         9, 900, 500, 90, 40, 70],
        ["202403", "E06000001", "E0701", "Unitary D", "submitted", "Unitary Authority", "Unitary Authority",
         1, 100, 60, 10, 4, 7],
        ["202503", "E06000001", "E0701", "Unitary D", "submitted", "Unitary Authority", "Unitary Authority",
         1, 150, 80, "Not set", 5, 8],
    ], columns=["year_ending", "ONS_code", "LA_LGF_code", "LA_name", "status", "LA_class", "LA_subclass",
                "RS_hous_net_exp", "RO4_tanps_tot_exp", "RO4_tanps_net_exp", "RO4_tabb_tot_exp",
                "RO4_taadm_tot_exp", "RO4_hra_tot_exp"]).to_csv(csv, index=False)
    df = read_ro_timeseries(csv, meta)
    assert list(df.columns) == TA_SPEND_COLUMNS
    assert set(df["la_code"]) == {"E06000001"}  # class total row dropped
    assert set(df["year"]) == {"2023-24", "2024-25"}
    assert set(df["line"]) == {"nightly_paid_self_contained", "bed_and_breakfast", "ta_admin"}
    row = df[(df["year"] == "2024-25") & (df["line"] == "nightly_paid_self_contained")]
    assert row["gbp"].iloc[0] == 150_000
    assert np.isnan(df[(df["year"] == "2024-25") & (df["line"] == "bed_and_breakfast")]["gbp"].iloc[0])
    net = read_ro_timeseries(csv, meta, measure="net_current_expenditure")
    assert set(net["line"]) == {"nightly_paid_self_contained"}
    explicit = read_ro_timeseries(csv, items={"RO4_tabb_tot_exp": "bed_and_breakfast"}, include_aggregates=True)
    assert "E06" in set(explicit["la_code"])


def test_harmonise_codes_with_repo_table():
    from pathlib import Path

    changes = mhclg.read_code_changes(Path(__file__).resolve().parents[1] / "config" / "code_changes.csv")
    assert changes["old_code"].is_unique
    assert changes["old_code"].str.match(r"^E0[6-9]\d{6}$").all()
    assert changes["new_code"].str.match(r"^E0[6-9]\d{6}$").all()
    df = pd.DataFrame([
        ("E07000190", "Taunton Deane", "2019-Q1", "total", 10.0, "a.ods"),  # two steps to Somerset
        ("E07000187", "Mendip", "2019-Q1", "total", 7.0, "a.ods"),
        ("E07000187", "Mendip", "2022-Q4", "total", 8.0, "b.ods"),
        ("E07000188", "Sedgemoor", "2022-Q4", "total", np.nan, "b.ods"),  # a gap stays a gap
        ("E09000001", "Borough A", "2022-Q4", "total", 3.0, "b.ods"),
    ], columns=TA_HOUSEHOLDS_COLUMNS)
    harmonised = mhclg.harmonise_codes(df, changes, "households")
    assert list(harmonised.columns) == TA_HOUSEHOLDS_COLUMNS
    out = harmonised.set_index(["la_code", "quarter"])
    assert out.loc[("E06000066", "2019-Q1"), "households"] == 17
    assert out.loc[("E06000066", "2019-Q1"), "la_name"] == "Somerset"
    assert np.isnan(out.loc[("E06000066", "2022-Q4"), "households"])
    assert out.loc[("E09000001", "2022-Q4"), "households"] == 3


def test_discover_raw_files(tmp_path):
    for name in ["Statutory_Homelessness_Detailed_Local_Authority_Data_202603.ods", "Detailed_LA_202403.xlsx",
                 "MDIS_Tables_202603.ods", "RO4_LA_Data_2024-25_data_by_LA.ods",
                 "Revenue_Outturn_time_series_data_v4.csv", "Revenue_Outturn_metadata.ods",
                 "Statutory_Homelessness_England_Time_Series_202603.ods",
                 "LAD25_RGN25_EN_LU.csv"]:
        (tmp_path / name).write_bytes(b"")
    groups = discover_raw_files(tmp_path)
    assert [p.name for p in groups["households"]] == [
        "Detailed_LA_202403.xlsx", "Statutory_Homelessness_Detailed_Local_Authority_Data_202603.ods"]
    assert [p.name for p in groups["spend"]] == ["RO4_LA_Data_2024-25_data_by_LA.ods"]
    assert len(groups["ro_timeseries"]) == len(groups["ro_metadata"]) == len(groups["lookup"]) == 1
    assert {p.name for p in groups["other"]} == {"MDIS_Tables_202603.ods",
                                                 "Statutory_Homelessness_England_Time_Series_202603.ods"}
